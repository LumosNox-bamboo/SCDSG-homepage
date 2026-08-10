#!/usr/bin/env python3
"""Build the one-page SCDSG Forum 2026 abstract preparation template."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "assets" / "docs" / "scdsg-forum-2026-abstract-template.docx"

# Resolved preset: compact_reference_guide. Named one-page-form overrides:
# 0.72/0.82-inch margins, tighter H1 spacing, and a 23-pt memo_masthead title.
BLUE = RGBColor(21, 94, 168)
DARK_BLUE = RGBColor(7, 29, 48)
MUTED = RGBColor(88, 104, 116)
GOLD = RGBColor(169, 130, 75)
LIGHT = "EDF3F7"


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_font(run, size=11, bold=False, color=DARK_BLUE, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_field(doc, label, prompt, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(f"{label}: "), bold=True, color=BLUE)
    set_font(p.add_run(prompt), color=MUTED, italic=True)


def add_abstract_block(doc, heading, prompt):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(heading), size=11, bold=True, color=DARK_BLUE)
    body = doc.add_paragraph()
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.line_spacing = 1.25
    set_font(body.add_run(f"[{prompt}]"), size=10.5, color=MUTED, italic=True)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)
section.header_distance = Inches(0.32)
section.footer_distance = Inches(0.32)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = DARK_BLUE
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for style_name, size, before, after, color in (
    ("Heading 1", 16, 12, 6, BLUE),
    ("Heading 2", 13, 9, 4, BLUE),
    ("Heading 3", 12, 7, 3, DARK_BLUE),
):
    style = doc.styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("SCDSG · YOUNG SCHOLARS FORUM 2026"), size=8.5, bold=True, color=MUTED)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("Preparation template · Submit online at scdsg-med.com/forum-2026/register/"), size=8, color=MUTED)

kicker = doc.add_paragraph()
kicker.paragraph_format.space_after = Pt(3)
set_font(kicker.add_run("CALL FOR ABSTRACTS"), size=9, bold=True, color=GOLD)

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(3)
set_font(title.add_run("Abstract Preparation Template"), size=23, bold=True, color=DARK_BLUE)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
set_font(subtitle.add_run("SCDSG Young Scholars Forum 2026 · English abstract · Maximum 300 words"), size=11, bold=True, color=BLUE)

note = doc.add_paragraph()
note.paragraph_format.left_indent = Inches(0.12)
note.paragraph_format.right_indent = Inches(0.12)
note.paragraph_format.space_before = Pt(4)
note.paragraph_format.space_after = Pt(8)
note.paragraph_format.line_spacing = 1.15
shade_paragraph(note, LIGHT)
set_font(note.add_run("Preparation aid only. "), size=9.5, bold=True, color=BLUE)
set_font(note.add_run("Do not upload this file. Copy the final title, abstract and keywords into the online submission form. All text entered in the abstract field, including section headings, must remain within 300 words."), size=9.5, color=DARK_BLUE)

doc.add_heading("1 · Contribution details", level=1)
add_field(doc, "Contribution title", "Replace with a concise English title")
add_field(doc, "Scientific track", "Select one track from the online submission form")
add_field(doc, "Presentation preference", "Oral presentation / Poster presentation / Either")
add_field(doc, "Keywords", "Enter 3–6 English keywords separated by semicolons", space_after=4)

doc.add_heading("2 · Abstract body", level=1)
intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(4)
intro.paragraph_format.line_spacing = 1.15
set_font(intro.add_run("Use the structure below when appropriate for your discipline. Keep statements specific and report results rather than planned analyses."), size=9.5, color=MUTED)

add_abstract_block(doc, "Background and objective", "State the clinical or scientific problem and the study objective")
add_abstract_block(doc, "Methods", "Describe design, sample or model, principal methods and analysis")
add_abstract_block(doc, "Results", "Report the main findings with quantitative results where available")
add_abstract_block(doc, "Conclusion", "State the conclusion supported by the reported results")

doc.add_heading("3 · Before submitting", level=1)
check = doc.add_paragraph()
check.paragraph_format.space_after = Pt(0)
check.paragraph_format.line_spacing = 1.2
set_font(check.add_run("Confirm: "), size=9.5, bold=True, color=BLUE)
set_font(check.add_run("English language · no more than 300 words · 3–6 keywords · presentation preference selected · CV prepared as PDF (maximum 10 MB) · optional supplementary figure prepared as JPG, PNG or WebP (maximum 20 MB)."), size=9.5, color=DARK_BLUE)

doc.core_properties.title = "SCDSG Forum 2026 Abstract Preparation Template"
doc.core_properties.subject = "Abstract preparation aid for the SCDSG Young Scholars Forum 2026"
doc.core_properties.author = "SCDSG"
doc.core_properties.keywords = "SCDSG, Forum 2026, abstract template"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
